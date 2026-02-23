from __future__ import annotations

import io
import math
import re
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any

import xlrd
from openpyxl import load_workbook

from app.schemas.signal_import_schema import SignalImportMetaSchema


_INTERNAL_TYPE_TO_DIRECTION: dict[str, str] = {
    "di": "DI",
    "do": "DO",
    "ai": "AI",
    "ao": "AO",
}


@dataclass(frozen=True)
class PackedSignalRole:
    suffix: str
    role_key: str


@dataclass(frozen=True)
class PackedSignalProfile:
    kind: str
    default_internal_type: str
    roles: tuple[PackedSignalRole, PackedSignalRole]


_PACKED_SIGNAL_PROFILES: dict[str, PackedSignalProfile] = {
    "dps": PackedSignalProfile(
        kind="dps",
        default_internal_type="di",
        roles=(
            PackedSignalRole(suffix=" / OPEN FB", role_key="open_fb"),
            PackedSignalRole(suffix=" / CLOSE FB", role_key="close_fb"),
        ),
    ),
    "dpc": PackedSignalProfile(
        kind="dpc",
        default_internal_type="do",
        roles=(
            PackedSignalRole(suffix=" / OPEN CMD", role_key="open_cmd"),
            PackedSignalRole(suffix=" / CLOSE CMD", role_key="close_cmd"),
        ),
    ),
}

_PACKED_SIGNAL_TYPE_ALIASES: dict[str, str] = {
    # DPC (double-point command)
    "dpc": "dpc",
    "dpc.": "dpc",
    "double point command": "dpc",
    "double-point command": "dpc",
    "double_point_command": "dpc",
    "doublepointcommand": "dpc",
    "dp command": "dpc",
    "double command": "dpc",
    # DPS (double-point status)
    "dps": "dps",
    "dps.": "dps",
    "double point status": "dps",
    "double-point status": "dps",
    "double_point_status": "dps",
    "doublepointstatus": "dps",
    "dp status": "dps",
    "double status": "dps",
}


@dataclass(frozen=True)
class ImportedSignalProjection:
    key: str
    name: str
    io_direction: str
    category: str | None
    signal_metadata: dict[str, Any]


@dataclass(frozen=True)
class ImportedSheetPayload:
    data: dict[str, Any]
    rows_count: int
    signals: list[ImportedSignalProjection]


class SignalSheetImportService:
    """Parse workbook bytes into normalized sheet payload + optional signal projection."""

    @staticmethod
    def parse_workbook(
        file_bytes: bytes,
        *,
        filename: str | None,
        metadata: SignalImportMetaSchema | None,
    ) -> ImportedSheetPayload:
        matrices = SignalSheetImportService._read_workbook(file_bytes, filename)
        sheets: list[dict[str, Any]] = []
        total_rows = 0

        for sheet_index, (sheet_name, matrix) in enumerate(matrices):
            headers = SignalSheetImportService._normalize_headers(matrix[0] if matrix else [])
            rows = SignalSheetImportService._rows_to_objects(matrix[1:], headers)
            total_rows += len(rows)
            sheets.append(
                {
                    "name": sheet_name,
                    "index": sheet_index,
                    "headers": headers,
                    "rows_count": len(rows),
                    "rows": rows,
                }
            )

        if not sheets:
            sheets.append(
                {
                    "name": "Sheet 1",
                    "index": 0,
                    "headers": [],
                    "rows_count": 0,
                    "rows": [],
                }
            )

        default_sheet_index = SignalSheetImportService._select_default_sheet_index(sheets, metadata)

        projected_signals = SignalSheetImportService._project_signals(
            sheets=sheets,
            default_sheet_index=default_sheet_index,
            filename=filename,
            metadata=metadata,
        )

        compact_sheets = SignalSheetImportService._project_sheet_data_columns(
            sheets=sheets,
            default_sheet_index=default_sheet_index,
            metadata=metadata,
        )

        sheet_data = {
            "version": 2,
            "sheet_count": len(compact_sheets),
            "default_sheet_index": default_sheet_index,
            "sheets": compact_sheets,
        }

        return ImportedSheetPayload(
            data=sheet_data,
            rows_count=total_rows,
            signals=projected_signals,
        )

    @staticmethod
    def _read_workbook(file_bytes: bytes, filename: str | None) -> list[tuple[str, list[list[Any]]]]:
        extension = (filename or "").lower().split(".")[-1] if filename and "." in filename else ""

        if extension == "xls":
            return SignalSheetImportService._read_xls(file_bytes)

        try:
            return SignalSheetImportService._read_xlsx(file_bytes)
        except Exception:
            return SignalSheetImportService._read_xls(file_bytes)

    @staticmethod
    def _read_xlsx(file_bytes: bytes) -> list[tuple[str, list[list[Any]]]]:
        workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
        result: list[tuple[str, list[list[Any]]]] = []
        for ws in workbook.worksheets:
            matrix: list[list[Any]] = []
            for row in ws.iter_rows(values_only=True):
                matrix.append(list(row))
            result.append((ws.title or "Sheet", matrix))
        return result

    @staticmethod
    def _read_xls(file_bytes: bytes) -> list[tuple[str, list[list[Any]]]]:
        workbook = xlrd.open_workbook(file_contents=file_bytes)
        result: list[tuple[str, list[list[Any]]]] = []
        for sheet in workbook.sheets():
            matrix: list[list[Any]] = []
            for row_index in range(sheet.nrows):
                matrix.append(sheet.row_values(row_index))
            result.append((sheet.name or "Sheet", matrix))
        return result

    @staticmethod
    def _normalize_headers(raw_headers: list[Any]) -> list[str]:
        normalized: list[str] = []
        seen: dict[str, int] = {}

        for index, raw in enumerate(raw_headers):
            value = SignalSheetImportService._stringify_cell(raw).strip()
            base = value if value else f"column_{index + 1}"
            suffix = seen.get(base, 0)
            seen[base] = suffix + 1
            header = base if suffix == 0 else f"{base}_{suffix + 1}"
            normalized.append(header)

        return normalized

    @staticmethod
    def _rows_to_objects(rows: list[list[Any]], headers: list[str]) -> list[dict[str, Any]]:
        if not headers:
            return []

        payload: list[dict[str, Any]] = []
        for row in rows:
            item: dict[str, Any] = {}
            non_empty = False
            for col_idx, header in enumerate(headers):
                value = SignalSheetImportService._normalize_cell(row[col_idx] if col_idx < len(row) else None)
                item[header] = value
                if value not in (None, ""):
                    non_empty = True
            if non_empty:
                payload.append(item)
        return payload

    @staticmethod
    def _normalize_cell(value: Any) -> Any:
        if value is None:
            return None
        if isinstance(value, datetime):
            return value.isoformat()
        if isinstance(value, date):
            return value.isoformat()
        if isinstance(value, bytes):
            try:
                return value.decode("utf-8")
            except UnicodeDecodeError:
                return value.decode("latin-1", errors="ignore")
        if isinstance(value, float) and math.isnan(value):
            return None
        return value

    @staticmethod
    def _select_default_sheet_index(
        sheets: list[dict[str, Any]],
        metadata: SignalImportMetaSchema | None,
    ) -> int:
        if metadata and metadata.sheet_name:
            for sheet in sheets:
                if sheet["name"] == metadata.sheet_name:
                    return int(sheet["index"])

        for sheet in sheets:
            if int(sheet.get("rows_count") or 0) > 0:
                return int(sheet["index"])

        return 0

    @staticmethod
    def _project_signals(
        *,
        sheets: list[dict[str, Any]],
        default_sheet_index: int,
        filename: str | None,
        metadata: SignalImportMetaSchema | None,
    ) -> list[ImportedSignalProjection]:
        selected = next((sheet for sheet in sheets if sheet["index"] == default_sheet_index), None)
        if not selected:
            return []

        rows = selected.get("rows") or []
        if not rows:
            return []

        headers = selected.get("headers") or []
        signal_name_column = SignalSheetImportService._pick_signal_name_column(headers)
        internal_type_column = SignalSheetImportService._pick_internal_type_column(headers, metadata)
        type_column = metadata.type_column if metadata and metadata.type_column in headers else None
        type_mapping = {
            str(key).strip().lower(): str(value).strip().lower()
            for key, value in (metadata.type_mapping if metadata else {}).items()
            if str(key).strip() and str(value).strip()
        }

        projections: list[ImportedSignalProjection] = []
        seen_keys: set[str] = set()
        selected_columns = [
            column
            for column in ((metadata.selected_columns if metadata else []) or [])
            if column in headers
        ]

        for row_index, row in enumerate(rows):
            type_info = SignalSheetImportService._resolve_type_info(
                row=row,
                internal_type_column=internal_type_column,
                type_column=type_column,
                type_mapping=type_mapping,
            )
            internal_type = type_info["internal_type"]
            direction = _INTERNAL_TYPE_TO_DIRECTION.get(internal_type)
            if not direction:
                continue

            display_name = (
                SignalSheetImportService._stringify_cell(row.get(signal_name_column)).strip()
                if signal_name_column
                else ""
            )
            if not display_name:
                display_name = f"Signal {row_index + 1}"

            category = None
            if type_column:
                category_raw = row.get(type_column)
                category_str = SignalSheetImportService._stringify_cell(category_raw).strip()
                category = category_str or None

            packed_profile = SignalSheetImportService._resolve_packed_profile(type_info)
            if packed_profile:
                source_group_id = f"{selected.get('index')}:{row_index}"
                terminal_values = SignalSheetImportService._split_terminal_values(
                    row=row,
                    terminal_column=metadata.terminal_column if metadata else None,
                    expected=len(packed_profile.roles),
                )
                for packed_pos, role in enumerate(packed_profile.roles):
                    packed_name = f"{display_name}{role.suffix}"
                    packed_key = SignalSheetImportService._make_unique_signal_key(packed_name, seen_keys)
                    row_payload = SignalSheetImportService._project_row_payload(
                        row=row,
                        selected_columns=selected_columns,
                    )
                    if (
                        terminal_values
                        and metadata
                        and metadata.terminal_column
                        and (
                            not selected_columns
                            or metadata.terminal_column in row_payload
                        )
                    ):
                        row_payload[metadata.terminal_column] = terminal_values[packed_pos]
                    signal_metadata = {
                        "source": "signal_sheet_import",
                        "source_filename": filename,
                        "sheet_name": selected.get("name"),
                        "sheet_index": selected.get("index"),
                        "row_index": row_index,
                        "row": row_payload,
                        "packed_group": {
                            "kind": packed_profile.kind,
                            "source_group_id": source_group_id,
                            "position": packed_pos,
                            "role": role.role_key,
                        },
                    }
                    projections.append(
                        ImportedSignalProjection(
                            key=packed_key,
                            name=packed_name,
                            io_direction=direction,
                            category=category,
                            signal_metadata=signal_metadata,
                        )
                    )
                continue

            key = SignalSheetImportService._make_unique_signal_key(display_name, seen_keys)
            signal_metadata = {
                "source": "signal_sheet_import",
                "source_filename": filename,
                "sheet_name": selected.get("name"),
                "sheet_index": selected.get("index"),
                "row_index": row_index,
                "row": SignalSheetImportService._project_row_payload(
                    row=row,
                    selected_columns=selected_columns,
                ),
            }

            projections.append(
                ImportedSignalProjection(
                    key=key,
                    name=display_name,
                    io_direction=direction,
                    category=category,
                    signal_metadata=signal_metadata,
                )
            )

        return projections

    @staticmethod
    def _pick_signal_name_column(headers: list[str]) -> str | None:
        heuristics = ("hmi", "name", "signal", "description")
        for header in headers:
            lowered = header.lower()
            if any(token in lowered for token in heuristics):
                return header
        return headers[0] if headers else None

    @staticmethod
    def _pick_internal_type_column(headers: list[str], metadata: SignalImportMetaSchema | None) -> str | None:
        if metadata and metadata.internal_type_column and metadata.internal_type_column in headers:
            return metadata.internal_type_column

        heuristics = ("internal_type", "internal type", "type", "тип")
        for header in headers:
            lowered = header.lower()
            if any(token in lowered for token in heuristics):
                return header
        return None

    @staticmethod
    def _resolve_type_info(
        *,
        row: dict[str, Any],
        internal_type_column: str | None,
        type_column: str | None,
        type_mapping: dict[str, str],
    ) -> dict[str, str]:
        raw_internal_type = ""
        raw_type = ""
        if internal_type_column:
            raw_internal_type = SignalSheetImportService._normalize_type_token(row.get(internal_type_column))
            if raw_internal_type in _INTERNAL_TYPE_TO_DIRECTION:
                return {"internal_type": raw_internal_type, "raw_type": raw_internal_type}

        if type_column:
            raw_type = SignalSheetImportService._normalize_type_token(row.get(type_column))
            mapped = type_mapping.get(raw_type)
            if mapped in _INTERNAL_TYPE_TO_DIRECTION:
                return {"internal_type": mapped, "raw_type": raw_type}

        packed_profile = SignalSheetImportService._resolve_packed_profile(
            {"internal_type": "", "raw_type": raw_type or raw_internal_type}
        )
        if packed_profile:
            return {"internal_type": packed_profile.default_internal_type, "raw_type": packed_profile.kind}

        return {"internal_type": "", "raw_type": raw_type or raw_internal_type}

    @staticmethod
    def _resolve_packed_profile(type_info: dict[str, str]) -> PackedSignalProfile | None:
        raw_type = (type_info.get("raw_type") or "").strip().lower()
        if not raw_type:
            return None
        normalized = _PACKED_SIGNAL_TYPE_ALIASES.get(raw_type, raw_type)
        return _PACKED_SIGNAL_PROFILES.get(normalized)

    @staticmethod
    def _normalize_type_token(value: Any) -> str:
        token = SignalSheetImportService._stringify_cell(value).strip().lower()
        token = re.sub(r"\s+", " ", token)
        return token

    @staticmethod
    def _make_unique_signal_key(display_name: str, seen_keys: set[str]) -> str:
        base_key = SignalSheetImportService._slugify(display_name)
        key = base_key
        suffix = 2
        while key in seen_keys:
            key = f"{base_key}_{suffix}"
            suffix += 1
        seen_keys.add(key)
        return key

    @staticmethod
    def _split_terminal_values(
        *,
        row: dict[str, Any],
        terminal_column: str | None,
        expected: int,
    ) -> list[str] | None:
        if expected <= 1 or not terminal_column or terminal_column not in row:
            return None

        raw = SignalSheetImportService._stringify_cell(row.get(terminal_column)).strip()
        if not raw:
            return None

        # Typical engineering sheets store paired terminals as "XT1:1 / XT1:2" or "XT1:1,XT1:2".
        parts = [part.strip() for part in re.split(r"\s*[/,;|]\s*", raw) if part and part.strip()]
        if len(parts) < expected:
            return None
        return parts[:expected]

    @staticmethod
    def _project_row_payload(*, row: dict[str, Any], selected_columns: list[str]) -> dict[str, Any]:
        if not selected_columns:
            return dict(row)
        return {column: row.get(column) for column in selected_columns}

    @staticmethod
    def _project_sheet_data_columns(
        *,
        sheets: list[dict[str, Any]],
        default_sheet_index: int,
        metadata: SignalImportMetaSchema | None,
    ) -> list[dict[str, Any]]:
        selected_columns = set(((metadata.selected_columns if metadata else []) or []))
        if not selected_columns:
            return [
                {
                    **sheet,
                    "headers": list(sheet.get("headers") or []),
                    "rows": [dict(row) for row in (sheet.get("rows") or []) if isinstance(row, dict)],
                }
                for sheet in sheets
            ]

        projected_sheets: list[dict[str, Any]] = []
        for sheet in sheets:
            headers = [str(header) for header in (sheet.get("headers") or [])]
            rows = [row for row in (sheet.get("rows") or []) if isinstance(row, dict)]
            if int(sheet.get("index") or 0) == int(default_sheet_index):
                filtered_headers = [header for header in headers if header in selected_columns]
                filtered_rows = [
                    {header: row.get(header) for header in filtered_headers}
                    for row in rows
                ]
                projected_sheets.append(
                    {
                        **sheet,
                        "headers": filtered_headers,
                        "rows": filtered_rows,
                    }
                )
                continue

            projected_sheets.append(
                {
                    **sheet,
                    "headers": headers,
                    "rows": [dict(row) for row in rows],
                }
            )
        return projected_sheets

    @staticmethod
    def _slugify(value: str) -> str:
        lowered = value.strip().lower()
        slug = re.sub(r"[^a-z0-9]+", "_", lowered).strip("_")
        return slug or "signal"

    @staticmethod
    def _stringify_cell(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, str):
            return value
        if isinstance(value, bool):
            return "true" if value else "false"
        return str(value)
