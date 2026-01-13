"""Signal snapshot import and lifecycle utilities."""
from __future__ import annotations

import hashlib
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable, Optional, TypedDict

from openpyxl import load_workbook
import xlrd  # type: ignore[import]

from sqlalchemy import Select, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.signal_snapshot import SignalSnapshot, SignalSnapshotStatus
from app.models.signal_snapshot_allocation import SignalSnapshotAllocation


CURRENT_SNAPSHOT_SCHEMA_VERSION = 2


class SnapshotSheetPayload(TypedDict):
    """Structured representation of a parsed worksheet."""

    name: str
    index: int
    headers: list[str]
    rows: list[dict[str, Any]]
    rows_count: int


class SnapshotWorkbookPayload(TypedDict):
    """Container for the workbook metadata persisted on the snapshot."""

    version: int
    sheet_count: int
    default_sheet_index: int
    sheets: list[SnapshotSheetPayload]


class InvalidSnapshotFileError(Exception):
    """Raised when the uploaded signal snapshot payload cannot be parsed."""


class SignalSnapshotNotFoundError(Exception):
    """Raised when a snapshot record cannot be located."""


class SnapshotLockedError(Exception):
    """Raised when a mutating action targets a locked snapshot."""


class SignalSnapshotService:
    """Coordinates persistence and validation for signal snapshots."""

    def __init__(self, db: AsyncSession):
        self.db = db

    async def import_workbook(self, workspace_id: int, payload: bytes, filename: str | None) -> SignalSnapshot:
        sheets = self._parse_workbook(payload, filename)
        workbook_payload = self._build_workbook_payload(sheets)
        snapshot = SignalSnapshot(
            workspace_id=workspace_id,
            status=SignalSnapshotStatus.DRAFT.value,
            source_filename=filename,
            source_hash=self._hash_payload(payload),
            rows_count=sum(sheet["rows_count"] for sheet in sheets),
            schema_version=CURRENT_SNAPSHOT_SCHEMA_VERSION,
            data=workbook_payload,
        )
        self.db.add(snapshot)
        await self.db.commit()
        await self.db.refresh(snapshot)
        return snapshot

    async def list_snapshots(self, workspace_id: int) -> list[SignalSnapshot]:
        stmt = (
            select(SignalSnapshot)
            .where(SignalSnapshot.workspace_id == workspace_id)
            .order_by(SignalSnapshot.created_at.desc())
        )
        result = await self.db.execute(stmt)
        return list(result.scalars().all())

    async def get_snapshot(self, snapshot_id: int) -> SignalSnapshot:
        snapshot = await self._get_snapshot(snapshot_id)
        if not snapshot:
            raise SignalSnapshotNotFoundError
        return snapshot

    async def delete_snapshot(self, snapshot_id: int) -> None:
        snapshot = await self._get_snapshot(snapshot_id, for_update=True)
        if not snapshot:
            raise SignalSnapshotNotFoundError
        if snapshot.is_locked():
            raise SnapshotLockedError
        await self.db.delete(snapshot)
        await self.db.commit()

    async def lock(self, snapshot_id: int) -> SignalSnapshot:
        """Lock snapshot by id via dedicated endpoint."""
        snapshot = await self._get_snapshot(snapshot_id, for_update=True)
        if not snapshot:
            raise SignalSnapshotNotFoundError
        self.lock_snapshot(snapshot)
        await self.db.commit()
        await self.db.refresh(snapshot)
        return snapshot

    async def get_allocation(self, snapshot_id: int) -> SignalSnapshotAllocation:
        snapshot = await self._get_snapshot(snapshot_id)
        if not snapshot:
            raise SignalSnapshotNotFoundError
        allocation = await self._get_allocation(snapshot_id)
        if allocation:
            return allocation
        allocation = SignalSnapshotAllocation(
            workspace_id=snapshot.workspace_id,
            signal_snapshot_id=snapshot.id,
            mapping=[],
        )
        self.db.add(allocation)
        await self.db.commit()
        await self.db.refresh(allocation)
        return allocation

    async def update_allocation(
        self,
        snapshot_id: int,
        mapping: list[dict[str, Any]],
    ) -> SignalSnapshotAllocation:
        snapshot = await self._get_snapshot(snapshot_id)
        if not snapshot:
            raise SignalSnapshotNotFoundError
        allocation = await self._get_allocation(snapshot_id)
        if allocation is None:
            allocation = SignalSnapshotAllocation(
                workspace_id=snapshot.workspace_id,
                signal_snapshot_id=snapshot.id,
                mapping=[],
            )
            self.db.add(allocation)
        allocation.mapping = mapping
        await self.db.commit()
        await self.db.refresh(allocation)
        return allocation

    @staticmethod
    def lock_snapshot(snapshot: SignalSnapshot) -> None:
        """Mark the snapshot as immutable for future operations."""
        if snapshot.status == SignalSnapshotStatus.LOCKED:
            return
        snapshot.status = SignalSnapshotStatus.LOCKED.value
        snapshot.locked_at = datetime.now(timezone.utc)

    async def _get_snapshot(
        self,
        snapshot_id: int,
        *,
        for_update: bool = False,
    ) -> SignalSnapshot | None:
        stmt: Select[tuple[SignalSnapshot]] = select(SignalSnapshot).where(SignalSnapshot.id == snapshot_id)
        if for_update:
            stmt = stmt.with_for_update()
        result = await self.db.execute(stmt)
        return result.scalar_one_or_none()

    async def _get_allocation(self, snapshot_id: int) -> SignalSnapshotAllocation | None:
        stmt: Select[tuple[SignalSnapshotAllocation]] = select(SignalSnapshotAllocation).where(
            SignalSnapshotAllocation.signal_snapshot_id == snapshot_id
        )
        result = await self.db.execute(stmt)
        return result.scalar_one_or_none()

    @staticmethod
    def _hash_payload(payload: bytes) -> str:
        return hashlib.sha256(payload).hexdigest()

    def _parse_workbook(self, payload: bytes, filename: str | None) -> list[SnapshotSheetPayload]:
        suffix = self._detect_suffix(filename)
        if suffix in {".xlsx", ".xlsm"}:
            sheets = self._parse_xlsx(payload)
        elif suffix == ".xls":
            sheets = self._parse_xls(payload)
        else:
            raise InvalidSnapshotFileError("Unsupported file format. Upload .xls, .xlsx, or .xlsm")

        if not any(sheet["rows_count"] for sheet in sheets):
            raise InvalidSnapshotFileError("Excel workbook contains no readable rows")

        return sheets

    @staticmethod
    def _detect_suffix(filename: str | None) -> str:
        if not filename:
            raise InvalidSnapshotFileError("Uploaded file must include a valid extension")
        return Path(filename).suffix.lower()

    def _parse_xlsx(self, payload: bytes) -> list[SnapshotSheetPayload]:
        try:
            workbook = load_workbook(filename=BytesIO(payload), read_only=True, data_only=True)
        except Exception as exc:  # pragma: no cover - delegated to library
            raise InvalidSnapshotFileError("Unable to read Excel workbook") from exc

        parsed: list[SnapshotSheetPayload] = []
        try:
            for index, sheet in enumerate(workbook.worksheets):
                header_row = self._normalize_header_row(
                    next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
                )
                rows = self._extract_rows_from_iterable(
                    header_row=header_row,
                    values=sheet.iter_rows(min_row=2, values_only=True),
                )
                parsed.append(
                    self._build_sheet_payload(
                        name=sheet.title,
                        index=index,
                        header_row=header_row,
                        rows=rows,
                    )
                )
        finally:
            workbook.close()

        return parsed

    def _parse_xls(self, payload: bytes) -> list[SnapshotSheetPayload]:
        try:
            book = xlrd.open_workbook(file_contents=payload, on_demand=True)
        except Exception as exc:  # pragma: no cover
            raise InvalidSnapshotFileError("Unable to read legacy Excel workbook") from exc

        parsed: list[SnapshotSheetPayload] = []
        try:
            for index, name in enumerate(book.sheet_names()):
                sheet = book.sheet_by_name(name)
                header_values = (
                    [sheet.cell_value(0, idx) for idx in range(sheet.ncols)]
                    if sheet.nrows
                    else []
                )
                rows_iter = (
                    tuple(
                        self._convert_xls_cell(book, sheet, row_idx, col_idx)
                        for col_idx in range(sheet.ncols)
                    )
                    for row_idx in range(1, sheet.nrows)
                )
                header_row = self._normalize_header_row(header_values)
                rows = self._extract_rows_from_iterable(
                    header_row=header_row,
                    values=rows_iter,
                )
                parsed.append(
                    self._build_sheet_payload(
                        name=name,
                        index=index,
                        header_row=header_row,
                        rows=rows,
                    )
                )
        finally:
            book.release_resources()

        return parsed

    @staticmethod
    def _build_workbook_payload(sheets: list[SnapshotSheetPayload]) -> SnapshotWorkbookPayload:
        default_index = SignalSnapshotService._default_sheet_index(sheets)
        return {
            "version": CURRENT_SNAPSHOT_SCHEMA_VERSION,
            "sheet_count": len(sheets),
            "default_sheet_index": default_index,
            "sheets": sheets,
        }

    @staticmethod
    def _default_sheet_index(sheets: list[SnapshotSheetPayload]) -> int:
        for sheet in sheets:
            if sheet["rows_count"] > 0:
                return sheet["index"]
        return sheets[0]["index"] if sheets else 0

    @staticmethod
    def _build_sheet_payload(
        *,
        name: str,
        index: int,
        header_row: list[tuple[int, str]],
        rows: list[dict[str, Any]],
    ) -> SnapshotSheetPayload:
        return {
            "name": name or f"Sheet {index + 1}",
            "index": index,
            "headers": [column_name for _, column_name in header_row],
            "rows": rows,
            "rows_count": len(rows),
        }

    @staticmethod
    def _normalize_header_row(raw_row: Optional[Iterable[Any]]) -> list[tuple[int, str]]:
        if raw_row is None:
            return []
        header: list[tuple[int, str]] = []
        for idx, value in enumerate(raw_row):
            if isinstance(value, str):
                normalized = value.strip()
            elif isinstance(value, datetime):
                normalized = value.isoformat()
            else:
                normalized = str(value).strip() if value is not None else ""
            if normalized:
                header.append((idx, normalized))
        return header

    def _extract_rows_from_iterable(
        self,
        *,
        header_row: list[tuple[int, str]],
        values: Iterable[Iterable[Optional[object]]],
    ) -> list[dict[str, Any]]:
        if not header_row:
            return []
        rows: list[dict[str, Any]] = []
        for raw_row in values:
            row_dict: dict[str, Any] = {}
            has_value = False
            row_list = list(raw_row)
            for column_index, column_name in header_row:
                cell_value = row_list[column_index] if column_index < len(row_list) else None
                normalized = self._normalize_cell(cell_value)
                row_dict[column_name] = normalized
                if normalized is not None and normalized != "":
                    has_value = True
            if has_value:
                rows.append(row_dict)
        return rows

    @staticmethod
    def _convert_xls_cell(book: xlrd.book.Book, sheet: xlrd.sheet.Sheet, row_idx: int, col_idx: int) -> Optional[object]:
        cell_type = sheet.cell_type(row_idx, col_idx)
        cell_value = sheet.cell_value(row_idx, col_idx)
        if cell_type == xlrd.XL_CELL_DATE:
            try:
                return datetime(*xlrd.xldate_as_tuple(cell_value, book.datemode))
            except Exception:  # pragma: no cover - fallback to raw value
                return cell_value
        return cell_value

    @staticmethod
    def _normalize_cell(value: Optional[object]) -> Optional[object]:
        if value is None:
            return None
        if isinstance(value, str):
            trimmed = value.strip()
            return trimmed or None
        if isinstance(value, datetime):
            return value.isoformat()
        return value
